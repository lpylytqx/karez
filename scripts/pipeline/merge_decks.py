#!/usr/bin/env python -u
# -*- coding: utf-8 -*-
"""把展示稿（10 页·画面型）与答辩稿（15 页·证据型）合并成一份完整 PPT。

合并思路 —— **按叙事顺序交错，而不是简单拼接**：
    ① 先看（展示稿 1-7）：封面 → 治水 → 岗位 → 四个画面 → 战斗 → 灾难 → 文化
    ② 再讲清（答辩稿 3-6）：坎儿井是主线 → 文化→机制映射 → 牧与工 → 乐与节
    ③ 技术（展示稿 8 + 答辩稿 8-11）：三句话 + 四个创新点
    ④ 数据（答辩稿 12-14 + 展示稿 9）：玩法深度 → 图表 → 交付状态 → 数字一览
    ⑤ 收尾（展示稿 10）

丢弃不要的：答辩稿的封面（与展示稿封面重复）、答辩稿的章节页（2/7，合并后不需要）、
答辩稿的收尾（与展示稿收尾重复）。

实现：python-pptx 没有跨文件复制页面的官方 API，所以手工搬 spTree 元素，
**并重映射图片关系**（a:blip 的 r:embed 指向源文件的关系 ID，必须在新文件里重建，
否则图片全部丢失）。这一步是整件事最容易出错的地方，所以搬完要逐个校验。
"""
from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _root import ROOT  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

D = ROOT / "deliverables"
SHOW = D / "坎儿井_作品展示.pptx"
DEF = D / "坎儿井_创新点答辩.pptx"
OUT = D / "坎儿井_完整版.pptx"

W, H = 12192000, 6858000   # 13.333 x 7.5 in

# 合并顺序：('S', 页码) 来自展示稿，('D', 页码) 来自答辩稿
ORDER = [("S", 1), ("S", 2), ("S", 3), ("S", 4), ("S", 5), ("S", 6), ("S", 7),
         ("D", 3), ("D", 4), ("D", 5), ("D", 6),
         ("S", 8),
         ("D", 8), ("D", 9), ("D", 10), ("D", 11),
         ("D", 12), ("D", 13), ("D", 14),
         ("S", 9), ("S", 10)]


def copy_slide(src_slide, dst_prs):
    """把一个 slide 的元素搬到新文稿，并修复图片关系。"""
    new = dst_prs.slides.add_slide(dst_prs.slide_layouts[6])
    # 清掉空白版式自带的占位符
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)

    # 按源顺序搬（保持 z-order：底图在前）
    for shp in src_slide.shapes:
        new.shapes._spTree.append(copy.deepcopy(shp._element))

    # ★ 重映射图片：a:blip 的 r:embed 是**源文件**里的关系 id，
    #   不重建关系的话，搬过来的图全部显示不出来
    n_pic = 0
    for el in new.shapes._spTree.iter():
        if el.tag != qn("a:blip"):
            continue
        rid = el.get(qn("r:embed"))
        if not rid:
            continue
        try:
            part = src_slide.part.related_part(rid)
            new_rid, _ = new.part.get_or_add_image_part(io.BytesIO(part.blob))
            el.set(qn("r:embed"), new_rid)
            n_pic += 1
        except Exception as e:
            print("      [警告] 图片关系重映射失败 rId=%s: %s" % (rid, e))

    # 讲者备注（答辩稿每页都有，合并后仍然有用）
    try:
        if src_slide.has_notes_slide and src_slide.notes_slide.notes_text_frame.text.strip():
            new.notes_slide.notes_text_frame.text = \
                src_slide.notes_slide.notes_text_frame.text
    except Exception:
        pass
    return new, n_pic


def main() -> None:
    show = Presentation(SHOW)
    defe = Presentation(DEF)
    print("  源：展示稿 %d 页 / 答辩稿 %d 页  尺寸 %s x %s / %s x %s"
          % (len(show.slides), len(defe.slides),
             show.slide_width, show.slide_height,
             defe.slide_width, defe.slide_height))
    if (show.slide_width, show.slide_height) != (defe.slide_width, defe.slide_height):
        print("  [X] 两份尺寸不同，不能合")
        return

    out = Presentation()
    out.slide_width, out.slide_height = W, H
    pics = 0
    for i, (tag, pg) in enumerate(ORDER, 1):
        src = (show if tag == "S" else defe).slides[pg - 1]
        _, n = copy_slide(src, out)
        pics += n
        if n:
            print("    第 %2d 页  <- %s 第 %d 页  图片 %d 张" % (i, tag, pg, n))
    out.save(OUT)
    print("\n  已生成 %s" % OUT.name)
    print("    %d 页（展示稿 10 页取 10，答辩稿 15 页取 11）" % len(out.slides))

    # ★ 校验：回读新文件，确认每页元素数与图片数都没丢
    back = Presentation(OUT)
    print("  回读校验：")
    for i, s in enumerate(back.slides, 1):
        blips = sum(1 for el in s.shapes._spTree.iter() if el.tag == qn("a:blip"))
        txts = sum(1 for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())
        print("    第 %2d 页：形状 %2d  图片 %2d  文本块 %2d" % (i, len(s.shapes), blips, txts))


if __name__ == "__main__":
    main()
