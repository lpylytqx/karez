#!/usr/bin/env python -u
# -*- coding: utf-8 -*-
"""两份 PPT 的收尾自检：几何越界 + 署名用字。

⚠ 这个动作本来在 pwsh 里用 `python -c` 跑，**静默失败了两次**（多行 Python 经
   PowerShell 转义会坏，输出又被打到 $null）。所以固化成文件。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _root import DELIVERABLES  # noqa: E402

from pptx import Presentation  # noqa: E402

DECKS = ("坎儿井_作品展示.pptx", "坎儿井_创新点答辩.pptx")


def main() -> None:
    print("  收尾自检：")
    for n in DECKS:
        p = Presentation(DELIVERABLES / n)
        W, H = p.slide_width, p.slide_height
        bad = 0
        txt = ""
        for s in p.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    txt += sh.text_frame.text + "\n"
                try:
                    if sh.left + (sh.width or 0) > W + 10000 or \
                       sh.top + (sh.height or 0) > H + 10000:
                        bad += 1
                except Exception:
                    pass
        wrong = re.findall(r"王焕(?!楷)", txt)
        print("    %-24s %2d 页  越界 %d %s   徐媛媛×%d  姓名用字 %s"
              % (n, len(p.slides), bad, "✓" if bad == 0 else "✗",
                 txt.count("徐媛媛"), "有错 ✗" if wrong else "无误 ✓"))
    # 裁图是否齐备
    c = DELIVERABLES / "assets" / "crops"
    fs = sorted(c.glob("*.png"))
    print("    裁图目录 %d 张 %s" % (len(fs), "✓" if len(fs) >= 9 else "✗"))


if __name__ == "__main__":
    main()
