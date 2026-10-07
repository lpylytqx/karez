#!/usr/bin/env python -u
# -*- coding: utf-8 -*-
"""核对两份 PPT 的署名与姓名用字，并同步到参赛材料。"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _root import DELIVERABLES, PACK  # noqa: E402

from pptx import Presentation  # noqa: E402


def main() -> None:
    print("  署名核对：")
    for n in ("坎儿井_作品展示.pptx", "坎儿井_创新点答辩.pptx"):
        p = Presentation(DELIVERABLES / n)
        txt = ""
        for s in p.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    txt += sh.text_frame.text + "\n"
        wrong = re.findall(r"王焕(?!楷)", txt)
        print("    %-24s 徐媛媛×%-2d 李中岩×%-2d 指导教师×%-2d "
              "王焕楷×%-2d  李云天×%-2d  %s"
              % (n, txt.count("徐媛媛"), txt.count("李中岩"),
                 txt.count("指导教师"), txt.count("王焕楷"),
                 txt.count("李云天"), "⚠ 有「王焕」未接楷" if wrong else "✓ 用字无误"))

    # 几何自检（不靠渲染图）
    for n in ("坎儿井_作品展示.pptx", "坎儿井_创新点答辩.pptx"):
        p = Presentation(DELIVERABLES / n)
        W, H = p.slide_width, p.slide_height
        bad = 0
        for s in p.slides:
            for sh in s.shapes:
                try:
                    if sh.left + (sh.width or 0) > W + 10000 or \
                       sh.top + (sh.height or 0) > H + 10000:
                        bad += 1
                except Exception:
                    pass
        print("    %-24s %d 页，越界元素 %d 个 %s"
              % (n, len(p.slides), bad, "✓" if bad == 0 else "✗"))

    for d in (PACK / "04-答辩PPT", PACK / "90-原始材料"):
        for n in ("坎儿井_作品展示.pptx", "坎儿井_创新点答辩.pptx"):
            shutil.copy2(DELIVERABLES / n, d / n)
    print("  两份 PPT 已同步到 04-答辩PPT/ 与 90-原始材料/")


if __name__ == "__main__":
    main()
