#!/usr/bin/env python -u
# -*- coding: utf-8 -*-
"""查合并稿里每页的页脚页码，确认统一了、且没有漏页。"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _root import ROOT  # noqa: E402

from pptx import Presentation  # noqa: E402

OUT = ROOT / "deliverables" / "坎儿井_完整版.pptx"
NUM = re.compile(r"^\s*(\d{1,2})\s*/\s*(\d{1,2})\s*$")


def main() -> None:
    p = Presentation(OUT)
    n = len(p.slides)
    print("  合并稿 %d 页。逐页页脚：" % n)
    miss = []
    for i, s in enumerate(p.slides, 1):
        hit = ""
        for sh in s.shapes:
            if sh.has_text_frame:
                m = NUM.match(sh.text_frame.text)
                if m:
                    hit = sh.text_frame.text.strip()
                    break
        if hit:
            ok = hit == "%02d / %d" % (i, n)
            print("    第 %2d 页  页脚 %-8s %s" % (i, hit, "✓" if ok else "✗ 对不上"))
            if not ok:
                miss.append(i)
        else:
            print("    第 %2d 页  页脚 （无）" % i)
            miss.append(i)
    print("\n  没有正确页码的页：%s" % (miss if miss else "无 ✓"))


if __name__ == "__main__":
    main()
